"""Wycena w efenix przez rdzeń Uruchomienia procesu na Atrapie efenix w prawdziwej przeglądarce (spec #741, #743, #746, #747).

Każdy test prowadzi ``run_process`` z opisem procesu ``efenix_process.PROCESS`` i oknem efenix na karcie Chromium, a
sprawdza to, co widać z zewnątrz: wynik, Komunikaty, liczby podsumowania, żądania, które dostała Atrapa efenix, jej
stan po przebiegu (ilości, data wykonania, załadowania Wyceny Emmetta, zmiany statusu) i Plik zbiorczy po przebiegu
(Znaczniki efenix czytane tą samą drogą co Przebieg rozliczeń, ``read_data_rows``, i części zip pliku). Przypisanie kolumn do pozycji
testy piszą literałami ze specyfikacji #741 (O→012, T→019, Z→010, AG→015). Wynik w kształcie karty Uruchomienia
sprawdza Skrypt dla maszyn (``tests/test_uruchom_proces.py``).
"""

from __future__ import annotations

from collections.abc import Callable, Iterator, Sequence
from contextlib import AbstractContextManager, contextmanager
from copy import copy
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path
import re
import socket
import sys
import zipfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection

from playwright.sync_api import Browser, Page
import pytest

from rozliczenia import efenix_browser, efenix_process
from rozliczenia.efenix_mark import MARK_RGB
from rozliczenia.efenix_pricing import PricingSummary
from rozliczenia.efenix_process import EfenixInputs
from rozliczenia.engine import run_settlements
from rozliczenia.process_run import Finished, Outcome, Stop, run_process
from rozliczenia.workbook_files import excel_lock_file, read_computed_rows, read_data_rows

from tests.fixtures.atrapa_efenix import (
    LIST_PAGE,
    LOAD_EMMETT,
    ORDER_PAGE,
    SAVE_QUANTITY,
    SEARCH,
    SET_DATE,
    AtrapaEfenix,
    AtrapaOrder,
    Changed,
    amount_of,
    atrapa_efenix,
)
from tests.fixtures.efenix import DAY, Order, write_collective_file
from tests.fixtures.interlocutor import Tata
from tests.fixtures.przegladarka import chromium, page_of
from tests.fixtures.settlement import PERIOD, write_mapping, write_template
from tests.fixtures.workbooks import active_worksheet, rewrite_parts

# Żądania, które niczego w efenix nie zmieniają: lista zleceń, szukanie i strona zlecenia.
READING = {("GET", LIST_PAGE), ("POST", SEARCH), ("GET", ORDER_PAGE)}
Priced = Outcome[PricingSummary]


@pytest.fixture(scope="module")
def browser() -> Iterator[Browser]:
    with chromium() as launched:
        yield launched


@pytest.fixture
def page(browser: Browser) -> Iterator[Page]:
    with page_of(browser) as opened:
        yield opened


def _price(tmp_path: Path, page: Page, url: str, orders: Sequence[Order], *, timeout_ms: int = 10_000) -> Priced:
    return _run(write_collective_file(tmp_path, orders), page, url, timeout_ms=timeout_ms)


def _run(
    source: Path, page: Page, url: str, *, timeout_ms: int = 10_000, meanwhile: Callable[[], None] = lambda: None
) -> Priced:
    """Uruchomienie Wyceny w efenix dla Pliku zbiorczego ``source``; ``meanwhile`` to to, co tata zrobił z plikiem, zanim
    proces skończył ostatnie zlecenie."""

    @contextmanager
    def open_efenix() -> Iterator[efenix_browser.BrowserEfenix]:
        with efenix_browser.window(page, url, timeout_ms=timeout_ms) as efenix:
            yield efenix
        meanwhile()

    return run_process(
        efenix_process.PROCESS,
        EfenixInputs(source, open_efenix),
        Tata(),
        view=lambda _inputs, summary: summary,
        metrics_path=source.parent.parent / "metryki.jsonl",
    )


def _figures(outcome: Priced) -> dict[str, int]:
    """Liczby podsumowania: zlecenia w zakresie, gotowe, błędy, ostrzeżenia Emmetta i pominięte."""

    assert isinstance(outcome.ending, Finished), outcome.ending
    assert outcome.view is not None
    summary = outcome.view
    return {
        "zakres": len(summary.results),
        "gotowe": summary.ready_count,
        "błędy": summary.error_count,
        "ostrzeżenia": summary.warning_count,
        "ze zdjęciami": summary.with_photos,
        "usterki": summary.defects,
    }


def _errors(outcome: Priced) -> list[tuple[str, str]]:
    """Kod i wiersz Pliku zbiorczego każdego błędu zlecenia i ostrzeżenia Emmetta."""

    assert outcome.view is not None
    return [(warning.code, warning.where.rpartition(", ")[2]) for warning in outcome.view.warnings]


def _what(outcome: Priced) -> list[str]:
    assert outcome.view is not None
    return [warning.what for warning in outcome.view.warnings]


def _pricing(atrapa: AtrapaEfenix) -> dict[str, tuple[dict[str, int], date | None, int]]:
    """Stan wyceny każdego zlecenia Atrapy: ilości, data wykonania i liczba załadowań Wyceny Emmetta."""

    orders, _logged_in = atrapa.state()
    return {order.number: (order.quantities, order.execution_date, order.emmett_loads) for order in orders}


def _stop(outcome: Priced) -> str:
    assert isinstance(outcome.ending, Stop), outcome.ending
    assert outcome.ending.kind == "stop"
    return outcome.ending.message.code


def _kept(outcome: Priced) -> bool:
    """Komunikat Zatrzymania mówi, że wycena zleceń sprzed niego została w efenix (``False``: że nic się nie zmieniło)."""

    assert isinstance(outcome.ending, Stop), outcome.ending
    what = outcome.ending.message.what
    kept, unchanged = "Wycena zleceń sprzed zatrzymania została w efenix." in what, "W efenix nic się nie zmieniło." in what
    assert kept != unchanged, what
    return kept


def _figures_of(orders: int, ready: int, warnings: int = 0, photos: int = 0, defects: int = 0) -> dict[str, int]:
    return {
        "zakres": orders,
        "gotowe": ready,
        "błędy": orders - ready - warnings,
        "ostrzeżenia": warnings,
        "ze zdjęciami": photos,
        "usterki": defects,
    }


def _au(quantities: dict[str, int]) -> float:
    """``Kwota`` (AU) zlecenia z ilościami ``quantities`` (kod pozycji → ilość) w cenach Atrapy."""

    return float(amount_of(quantities))


def test_zlecenie_z_wycena_zgodna_z_plikiem_jest_gotowe_a_efenix_dostaje_tylko_szukanie_i_otwarcie_zlecenia(
    tmp_path: Path, page: Page
) -> None:
    """Wycena już jest, każda ilość i kwota zgadzają się z plikiem: proces niczego nie zmienia, także pustej daty."""

    orders = [
        Order("#5001", cells={"O": 1, "T": 2}, amount=_au({"012": 1, "019": 2}), quantities={"012": 1, "019": 2}),
        Order("#5002", cells={"AG": 3}, amount=_au({"015": 3}), quantities={"015": 3}, emmett={"015": 1}),
        Order("5003", cells={"Z": 1}, amount=_au({"010": 1}), quantities={"010": 1}),
    ]
    with atrapa_efenix(
        [orders[0].in_efenix("901"), replace(orders[1].in_efenix("902"), execution_date=DAY), orders[2].in_efenix("903")]
    ) as atrapa:
        before = atrapa.state()
        outcome = _price(tmp_path, page, atrapa.url, orders)
        after = atrapa.state()
        requested = atrapa.requested()

    assert (outcome.result, _figures(outcome), _errors(outcome)) == ("OK", _figures_of(3, 3), [])
    assert after == before
    assert set(requested) <= READING
    assert requested.count(("GET", ORDER_PAGE)) == 3


@pytest.mark.parametrize(
    ("in_efenix", "positions"),
    [
        pytest.param({"012": 1, "019": 1}, "pozycja INDEKS_019 (kolumna T)", id="inna-ilość"),
        pytest.param({"012": 1, "019": 2, "005": 1}, "pozycja INDEKS_005 (bez kolumny w Pliku zbiorczym)", id="pozycja-bez-kolumny"),
        pytest.param({"012": 1, "019": 2}, None, id="ilości-zgodne-kwota-inna"),
    ],
)
def test_wycena_ktora_juz_jest_i_rozni_sie_od_pliku_to_blad_z_kwotami_bez_zmian_i_bez_emmetta(
    tmp_path: Path, page: Page, in_efenix: dict[str, int], positions: str | None
) -> None:
    au = _au({"012": 1, "019": 2}) + (0 if positions else 1)
    order = Order("#5101", cells={"O": 1, "T": 2}, amount=au, quantities=in_efenix, emmett={"012": 1, "019": 2})
    with atrapa_efenix([order.in_efenix("901")]) as atrapa:
        before = atrapa.state()
        outcome = _price(tmp_path, page, atrapa.url, [order])
        after = atrapa.state()
        requested = atrapa.requested()

    assert (outcome.result, _figures(outcome), _errors(outcome)) == ("WYMAGA_SPRAWDZENIA", _figures_of(1, 0), [("EFENIX_WYCENA_INNA", "wiersz 18")])
    [what] = _what(outcome)
    assert f"kwota w efenix {amount_of(in_efenix):.2f} zł".replace(".", ",") in what
    assert f"Kwota (AU) {Decimal(str(au)):.2f} zł".replace(".", ",") in what
    assert (positions is None) == ("Inne ilości" not in what)
    assert positions is None or positions in what
    assert (after, set(requested) <= READING) == (before, True)


def test_pusta_wycena_dostaje_date_z_data_wycene_emmetta_i_ilosci_z_pliku_w_pustych_pozycjach_a_kwota_rowna_au_daje_gotowe(
    tmp_path: Path, page: Page
) -> None:
    """Emmett wpisał O (012) zgodnie z plikiem; proces dopisuje T (019) i AG (015), a Z z 0 w pliku zostaje puste.

    Data wykonania jest ustawiana z ``DATA`` tylko w zleceniu, które jej nie miało.
    """

    filled = {"012": 1, "019": 2, "015": 1}
    orders = [
        Order("#5201", day=date(2026, 9, 3), cells={"O": 1, "T": 2, "AG": 1, "Z": 0}, amount=_au(filled), emmett={"012": 1}),
        # AU z ułamkiem binarnym, jaki zostawia suma Excela: równa kwocie efenix co do grosza.
        Order("#5202", cells={"O": 2}, amount=_au({"012": 2}) + 1e-9, emmett={"012": 2}),
    ]
    with atrapa_efenix([orders[0].in_efenix("901"), replace(orders[1].in_efenix("902"), execution_date=DAY)]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, orders)
        pricing = _pricing(atrapa)
        requested = atrapa.requested()
        changes = atrapa.status_changes()

    assert (outcome.result, _figures(outcome), _errors(outcome)) == ("OK", _figures_of(2, 2), [])
    assert pricing == {"5201": (filled, date(2026, 9, 3), 1), "5202": ({"012": 2}, DAY, 1)}
    assert (requested.count(("POST", SET_DATE)), requested.count(("POST", SAVE_QUANTITY)), changes) == (1, 2, 0)


def test_kwota_efenix_inna_niz_au_po_dopisaniu_ilosci_to_blad_z_obiema_kwotami(tmp_path: Path, page: Page) -> None:
    """Kwoty różnią się o 40 groszy: równość jest co do grosza, nie co do złotówki."""

    order = Order("#5301", cells={"O": 1, "T": 2}, amount=_au({"012": 1, "019": 2}) + 0.4, emmett={"012": 1})
    with atrapa_efenix([order.in_efenix("901")]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, [order])
        pricing = _pricing(atrapa)

    assert (outcome.result, _figures(outcome), _errors(outcome)) == ("WYMAGA_SPRAWDZENIA", _figures_of(1, 0), [("EFENIX_INNA_KWOTA", "wiersz 18")])
    assert pricing == {"5301": ({"012": 1, "019": 2}, DAY, 1)}
    [what] = _what(outcome)
    efenix_amount, au = amount_of({"012": 1, "019": 2}), amount_of({"012": 1, "019": 2}) + Decimal("0.40")
    assert f"kwota w efenix to {efenix_amount:.2f} zł, a Kwota (AU) {au:.2f} zł".replace(".", ",") in what


@pytest.mark.parametrize(
    ("day", "result", "errors", "pricing"),
    [
        pytest.param(date(2026, 9, 1), "OK", [], ({"012": 1}, date(2026, 9, 1), 1), id="pierwszy-dzień-poprzedniego-miesiąca"),
        pytest.param(date(2026, 10, 12), "OK", [], ({"012": 1}, date(2026, 10, 12), 1), id="dziś"),
        pytest.param(
            date(2026, 8, 31), "WYMAGA_SPRAWDZENIA", [("EFENIX_DATA_POZA_ZAKRESEM", "wiersz 18")], ({}, None, 0), id="przed-poprzednim-miesiącem"
        ),
        pytest.param(date(2026, 10, 13), "WYMAGA_SPRAWDZENIA", [("EFENIX_DATA_POZA_ZAKRESEM", "wiersz 18")], ({}, None, 0), id="jutro"),
    ],
)
def test_dzien_z_data_spoza_zakresu_kalendarza_efenix_to_blad_bez_zapisu_daty_bez_emmetta_i_bez_ilosci(
    tmp_path: Path,
    page: Page,
    day: date,
    result: str,
    errors: list[tuple[str, str]],
    pricing: tuple[dict[str, int], date | None, int],
) -> None:
    """Kalendarz Atrapy ma za dziś 12 października 2026, więc efenix przyjmuje daty wykonania od 1 września do 12 października."""

    order = Order("#5601", day=day, cells={"O": 1}, amount=_au({"012": 1}))
    with atrapa_efenix([order.in_efenix("901")], today=date(2026, 10, 12)) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, [order])
        after = _pricing(atrapa)
        requested = atrapa.requested()

    assert (outcome.result, _errors(outcome), after) == (result, errors, {"5601": pricing})
    assert requested.count(("POST", SET_DATE)) == (1 if result == "OK" else 0)


def test_data_wykonania_ktorej_efenix_nie_zapisal_po_potwierdz_to_blad_bez_emmetta_i_bez_ilosci(tmp_path: Path, page: Page) -> None:
    """Skrypt strony wpisuje dzień do ``#data_install`` po każdej odpowiedzi, także gdy serwer daty nie zapisał."""

    order = Order("#5701", cells={"O": 1}, amount=_au({"012": 1}))
    with atrapa_efenix([order.in_efenix("901")], refuse_dates=True) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, [order])
        pricing = _pricing(atrapa)
        requested = atrapa.requested()

    assert (outcome.result, _errors(outcome), pricing) == ("WYMAGA_SPRAWDZENIA", [("EFENIX_DATA_NIEZAPISANA", "wiersz 18")], {"5701": ({}, None, 0)})
    assert (requested.count(("POST", SET_DATE)), requested.count(("POST", LOAD_EMMETT)), requested.count(("POST", SAVE_QUANTITY))) == (1, 0, 0)


@pytest.mark.parametrize(
    ("emmett", "conflict"),
    [
        pytest.param({"012": 2}, "pozycja INDEKS_012 (kolumna O): Emmett 2, plik 1", id="inna-ilość"),
        pytest.param({"012": 1, "013": 1}, "pozycja INDEKS_013 (kolumna P): Emmett 1, plik 0", id="0-w-pliku"),
        pytest.param({"012": 1, "020": 1}, "pozycja INDEKS_020 (bez kolumny w Pliku zbiorczym): Emmett 1, plik 0", id="pozycja-bez-kolumny"),
    ],
)
def test_ilosc_emmetta_inna_niz_w_pliku_to_ostrzezenie_z_pozycja_a_proces_niczego_nie_dopisuje(
    tmp_path: Path, page: Page, emmett: dict[str, int], conflict: str
) -> None:
    order = Order("#5401", cells={"O": 1, "P": 0, "T": 2}, amount=_au({"012": 1, "019": 2}), emmett=emmett)
    with atrapa_efenix([order.in_efenix("901")]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, [order])
        pricing = _pricing(atrapa)
        requested = atrapa.requested()

    assert (outcome.result, _figures(outcome), _errors(outcome)) == (
        "WYMAGA_SPRAWDZENIA",
        _figures_of(1, 0, warnings=1),
        [("EFENIX_SPRZECZNOSC_Z_EMMETTEM", "wiersz 18")],
    )
    assert conflict in _what(outcome)[0]
    assert (pricing, ("POST", SAVE_QUANTITY) in requested) == ({"5401": (emmett, DAY, 1)}, False)


@pytest.mark.parametrize(
    ("cells", "error"),
    [
        pytest.param({"O": 1, "N": 1}, "EFENIX_KOLUMNA_BEZ_POZYCJI", id="kolumna-N-bez-pary"),
        pytest.param({"AR": 2}, "EFENIX_KOLUMNA_BEZ_POZYCJI", id="kolumna-AR-bez-pary"),
        pytest.param({"O": "tak"}, "EFENIX_ILOSC_NIECALKOWITA", id="ilość-słowem"),
        pytest.param({"O": -1}, "EFENIX_ILOSC_NIECALKOWITA", id="ilość-ujemna"),
    ],
)
def test_ilosc_ktorej_nie_da_sie_przepisac_do_efenix_to_blad_zlecenia_zanim_proces_go_szuka(
    tmp_path: Path, page: Page, cells: dict[str, int | str], error: str
) -> None:
    orders = [Order("#5501", cells=cells, amount=100.0), Order("#5502")]
    with atrapa_efenix([order.in_efenix(f"90{index}") for index, order in enumerate(orders)]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, orders)
        searched = atrapa.searched()

    assert (_figures(outcome), _errors(outcome), searched) == (_figures_of(2, 1), [(error, "wiersz 18")], ["5502"])


def test_szukanie_ktore_lista_efenix_wysyla_sama_po_powrocie_ze_strony_zlecenia_nie_zmienia_wynikow_szukania(
    tmp_path: Path, page: Page
) -> None:
    """Efenix odpowiada wolno, więc szukanie poprzedniego numeru, które lista wysyła sama po wczytaniu, trwa jeszcze,
    gdy okno szuka następnego zlecenia, i jego odpowiedź przychodzi pierwsza."""

    orders = [Order("#5101"), Order("#5102"), Order("#5103")]
    with atrapa_efenix([order.in_efenix(f"91{index}") for index, order in enumerate(orders)], answer_delay_s=0.3) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, orders)
        searched, by_list = atrapa.searched(), atrapa.searched_by_list()

    assert (outcome.result, _figures(outcome), _errors(outcome)) == ("OK", _figures_of(3, 3), [])
    assert (searched, by_list) == (["5101", "5102", "5103"], ["5101", "5102"])


def test_zlecenia_ze_zdjeciami_i_usterki_sa_pominiete_i_policzone_a_wiersz_z_samymi_zerami_jest_w_zakresie(
    tmp_path: Path, page: Page
) -> None:
    photo_columns = ["AB", "AC", "AD", "AE", "AF", "AJ", "AK", "AL", "AM", "AN", "AO", "AP", "AS"]
    orders = [
        *(Order(f"#60{index:02}", cells={column: 1}) for index, column in enumerate(photo_columns)),
        Order("#6101", cells={"AJ": "tak"}),
        Order("#6201", cells={"AI": 1}),
        Order("#6301", cells=dict.fromkeys([*photo_columns, "AI"], 0)),
        Order("#6302", cells=dict.fromkeys([*photo_columns, "AI"], "0")),
        Order("#6401", cells={"H": None}),
    ]
    with atrapa_efenix([order.in_efenix(f"9{index:02}") for index, order in enumerate(orders)]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, orders)
        searched = atrapa.searched()

    assert (outcome.result, _figures(outcome), searched) == ("OK", _figures_of(2, 2, photos=14, defects=1), ["6301", "6302"])


def test_zlecenia_ida_po_wykonawcach_alfabetycznie_a_u_kazdego_od_najstarszej_daty(tmp_path: Path, page: Page) -> None:
    orders = [
        Order("#7001", worker="zenon.testowy", day=date(2026, 9, 8)),
        Order("#7002", worker="adrian.testowy", day=date(2026, 9, 12)),
        Order("#7003", worker="Bogdan.testowy", day=date(2026, 9, 10)),
        Order("#7004", worker="adrian.testowy", day=date(2026, 9, 9)),
        Order("#7005", worker="zenon.testowy", day=date(2026, 9, 7)),
    ]
    with atrapa_efenix([order.in_efenix(f"90{index}") for index, order in enumerate(orders)]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, orders)
        searched = atrapa.searched()

    assert (outcome.result, searched) == ("OK", ["7004", "7002", "7003", "7005", "7001"])


def test_kazda_niezgodnosc_z_efenix_to_blad_zlecenia_z_powodem_a_proces_idzie_dalej(tmp_path: Path, page: Page) -> None:
    orders = [
        Order(""),  # wiersz 18: bez NUMER ZL.
        Order("#8002", day=None),  # 19: bez daty w DATA
        Order("#8003"),  # 20: efenix go nie zna
        Order("#8004"),  # 21: dwa zlecenia o tym numerze
        Order("#8005"),  # 22: inny status
        Order("#8006"),  # 23: inny realizator
        Order("#8007"),  # 24: inny adres
        Order("#8008"),  # 25: inna data wykonania
        Order("#8009", cells={"O": 1}, amount=_au({"012": 1}), quantities={"012": 1}),  # 26: gotowe, po wszystkich błędach
    ]
    efenix = [
        orders[1].in_efenix("902"),
        replace(orders[2].in_efenix("903"), number="80031"),
        orders[3].in_efenix("904"),
        orders[3].in_efenix("914"),
        replace(orders[4].in_efenix("905"), status="Do akceptacji"),
        replace(orders[5].in_efenix("906"), realizator="zenon.testowy"),
        replace(orders[6].in_efenix("907"), address="Miasto-Dzielnica ul.Testowa 14 Lokal 3"),
        replace(orders[7].in_efenix("908"), execution_date=date(2026, 9, 10)),
        replace(orders[8].in_efenix("909"), execution_date=DAY),
    ]
    with atrapa_efenix(efenix) as atrapa:
        before = atrapa.state()
        outcome = _price(tmp_path, page, atrapa.url, orders)
        after = atrapa.state()

    assert (outcome.result, _figures(outcome)) == ("WYMAGA_SPRAWDZENIA", _figures_of(9, 1))
    # Wiersz bez daty idzie na koniec swojego Wykonawcy, więc błędy są w kolejności wierszy dopiero po sortowaniu.
    assert sorted(_errors(outcome), key=lambda error: error[1]) == [
        ("EFENIX_BRAK_NUMERU", "wiersz 18"),
        ("EFENIX_BRAK_DATY", "wiersz 19"),
        ("EFENIX_NIEZNANY_NUMER", "wiersz 20"),
        ("EFENIX_KILKA_ZLECEN", "wiersz 21"),
        ("EFENIX_INNY_STATUS", "wiersz 22"),
        ("EFENIX_INNY_REALIZATOR", "wiersz 23"),
        ("EFENIX_INNY_ADRES", "wiersz 24"),
        ("EFENIX_INNA_DATA", "wiersz 25"),
    ]
    assert after == before


@pytest.mark.parametrize(
    ("efenix_address", "found"),
    [
        pytest.param("Poznań-Wilda ul.Kowalskiego Jana 12A Lokal 34", True, id="słowa-w-innej-kolejności-i-Lokal"),
        pytest.param("Poznań-Wilda ul.Jana Kowalskiego 12A Lokal 35", False, id="inny-lokal"),
        pytest.param("Poznań-Wilda ul.Jana Kowalskiego 12B Lokal 34", False, id="inny-numer"),
        pytest.param("Poznań-Wilda ul.Jana Kowalskiej 12A Lokal 34", False, id="inne-słowo"),
        pytest.param("Poznań-Wilda os.Jana Kowalskiego 12A Lokal 34", False, id="inny-rodzaj"),
    ],
)
def test_dane_adresowe_efenix_to_adres_z_pliku_w_dowolnej_kolejnosci_slow_z_lokalem_i_tylko_wtedy(
    tmp_path: Path, page: Page, efenix_address: str, found: bool
) -> None:
    order = Order("#9001", address="UL. JANA KOWALSKIEGO 12A/34", efenix_address=efenix_address)
    with atrapa_efenix([order.in_efenix("901")]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, [order])

    assert _errors(outcome) == ([] if found else [("EFENIX_INNY_ADRES", "wiersz 18")])


def test_plik_bez_kwot_au_policzonych_przez_excel_zatrzymuje_wycene_zanim_otworzy_efenix(tmp_path: Path, page: Page) -> None:
    orders = [Order("#9101"), Order("#9102", amount=None)]
    with atrapa_efenix([order.in_efenix(f"90{index}") for index, order in enumerate(orders)]) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, orders)
        requested = atrapa.requested()

    assert isinstance(outcome.ending, Stop)
    assert (_stop(outcome), outcome.ending.message.where.rpartition(", ")[2], requested) == ("EFENIX_KWOTY_NIEPOLICZONE", "wiersz 19", [])
    assert "Otwórz i zapisz Plik zbiorczy w Excelu" in outcome.ending.message.action


@pytest.mark.parametrize(
    ("expire_after", "searched", "kept"),
    [
        pytest.param(None, [], False, id="przed-szukaniem"),
        pytest.param((ORDER_PAGE, 0), ["9201"], False, id="przy-otwarciu-pierwszego-zlecenia"),
        pytest.param((SEARCH, 1), ["9201", "9202"], True, id="przy-szukaniu"),
        # Strona pierwszego zlecenia otwiera się trzy razy: przy otwarciu, po dacie wykonania i po Wycenie Emmetta.
        pytest.param((ORDER_PAGE, 3), ["9201", "9202"], True, id="przy-otwarciu-zlecenia"),
        pytest.param((SET_DATE, 1), ["9201", "9202"], True, id="przy-dacie-wykonania"),
        pytest.param((LOAD_EMMETT, 1), ["9201", "9202"], True, id="przy-wycenie-emmetta"),
        pytest.param((SAVE_QUANTITY, 1), ["9201", "9202"], True, id="przy-zapisie-ilości"),
    ],
)
def test_strona_logowania_przed_szukaniem_albo_w_trakcie_zatrzymuje_wycene_a_komunikat_mowi_czy_cos_zmienila(
    tmp_path: Path, page: Page, expire_after: tuple[str, int] | None, searched: list[str], kept: bool
) -> None:
    orders = [Order(f"#920{index}", cells={"O": 1}, amount=_au({"012": 1})) for index in (1, 2, 3)]
    with atrapa_efenix([order.in_efenix(f"90{index}") for index, order in enumerate(orders)], expire_after=expire_after) as atrapa:
        if expire_after is None:
            atrapa.expire_session()
        outcome = _price(tmp_path, page, atrapa.url, orders)

        assert (_stop(outcome), atrapa.searched(), _kept(outcome)) == ("EFENIX_WYLOGOWANY", searched, kept)


@pytest.mark.parametrize("changed", list(Changed))
def test_inna_budowa_tabeli_wynikow_albo_strony_zlecenia_zatrzymuje_wycene(tmp_path: Path, page: Page, changed: Changed) -> None:
    """Zatrzymanie przychodzi w czasie okna (tu 3 s), także przy elemencie obecnym, ale niewidocznym.

    Tylko kwota, która się nie odświeża, zatrzymuje wycenę po zmianie w efenix (Wycena Emmetta i ilość).
    """

    order = Order("#9501", cells={"O": 1}, amount=_au({"012": 1}))
    with atrapa_efenix([replace(order.in_efenix("901"), execution_date=DAY)], changed=changed) as atrapa:
        outcome = _price(tmp_path, page, atrapa.url, [order], timeout_ms=3_000)
        changes = atrapa.status_changes()

    assert (_stop(outcome), changes, _kept(outcome)) == ("EFENIX_ZMIENIONY", 0, changed is Changed.STALE_AMOUNT)


def test_strona_bez_listy_zlecen_i_adres_bez_efenix_zatrzymuja_wycene(tmp_path: Path, page: Page) -> None:
    with socket.socket() as unused:
        unused.bind(("127.0.0.1", 0))
        closed = f"http://127.0.0.1:{unused.getsockname()[1]}/{LIST_PAGE}"
    with atrapa_efenix([Order("#9301").in_efenix("901")]) as atrapa:
        changed = _price(tmp_path / "zmieniony", page, atrapa.url.replace(LIST_PAGE, "inna_strona.php"), [Order("#9301")])
    unreachable = _price(tmp_path / "niedostepny", page, closed, [Order("#9301")])

    assert (_stop(changed), _stop(unreachable), _kept(changed), _kept(unreachable)) == ("EFENIX_ZMIENIONY", "EFENIX_NIEDOSTEPNY", False, False)


# --- Znacznik efenix w Pliku zbiorczym (#747) ------------------------------------------------------------------------

YELLOW = "FFFF00"
CHANGED_PARTS = {"xl/worksheets/sheet1.xml", "xl/styles.xml"}


def _parts(path: Path) -> list[tuple[str, bytes]]:
    with zipfile.ZipFile(path) as archive:
        return [(info.filename, archive.read(info)) for info in archive.infolist()]


def _address_fills(source: Path) -> dict[int, str | None]:
    """RGB jednolitego wypełnienia ``ADRES`` każdego Wiersza danych tak, jak czyta je Przebieg rozliczeń (``None``: bez)."""

    fills: dict[int, str | None] = {}
    for row in read_data_rows(source, despite_excel_lock=True).rows:
        fill = row.cells[1].highlights.fill
        solid = isinstance(fill, PatternFill) and fill.patternType == "solid"
        fills[row.number] = str(fill.fgColor.rgb)[-6:].upper() if solid and isinstance(fill, PatternFill) else None
    return fills


def _marked(source: Path) -> list[int]:
    return [row for row, rgb in _address_fills(source).items() if rgb == MARK_RGB]


def _address_looks(source: Path) -> dict[int, tuple[Font, Border, Protection, Alignment, str]]:
    """Wygląd ``ADRES`` każdego Wiersza danych poza wypełnieniem: czcionka, obramowanie, ochrona, wyrównanie, format."""

    workbook = load_workbook(source)
    try:
        sheet = active_worksheet(workbook)
        return {
            row: (copy(cell.font), copy(cell.border), copy(cell.protection), copy(cell.alignment), cell.number_format)
            for row in range(18, sheet.max_row + 1)
            if (cell := sheet[f"B{row}"]).value is not None
        }
    finally:
        workbook.close()


def _amounts(source: Path) -> dict[int, object]:
    """``Kwota`` (AU) każdego Wiersza danych tak, jak policzył ją Excel."""

    return {row.number: row.value("AU") for row in read_computed_rows(source).rows}


def _without_address_styles(xml: bytes, rows: Sequence[int]) -> bytes:
    """XML arkusza bez atrybutu ``s`` komórek ``ADRES`` wierszy ``rows``."""

    def neutral(found: re.Match[bytes]) -> bytes:
        return found.group(1) if int(found.group(2)) in rows else found.group(0)

    return re.sub(rb'(<c r="B(\d+)")(?: s="\d+")?', neutral, xml)


STYLE_LISTS = {b"fills": b"fill", b"cellXfs": b"xf"}


def _style_list(styles: bytes, tag: bytes) -> re.Match[bytes]:
    found = re.search(rb"<" + tag + rb'\b[^>]*?count="(\d+)"[^>]*>(.*?)</' + tag + rb">", styles, re.S)
    assert found is not None, tag
    return found


def _style_counts(styles: bytes) -> dict[bytes, int]:
    """Liczba wpisów każdej listy ``fills`` i ``cellXfs`` w ``xl/styles.xml``, zgodna z jej atrybutem ``count``."""

    counts = {}
    for tag, entry in STYLE_LISTS.items():
        found = _style_list(styles, tag)
        counts[tag] = len(re.findall(rb"<" + entry + rb"[\s>/]", found.group(2)))
        assert int(found.group(1)) == counts[tag], tag
    return counts


def _without_appended_styles(written: bytes, original: bytes) -> bytes:
    """``xl/styles.xml`` ``written`` z listami ``fills`` i ``cellXfs`` z ``original``, gdy ``written`` tylko dopisał do
    nich wpisy na końcu."""

    for tag in STYLE_LISTS:
        before, after = _style_list(original, tag), _style_list(written, tag)
        assert after.group(2).startswith(before.group(2)), tag
        written = written[: after.start()] + before.group(0) + written[after.end() :]
    return written


def _marking_orders() -> list[Order]:
    """Wiersze 18 i 19 gotowe (wycena już jest; dopisane ilości, ``ADRES`` bez stylu), 20 błąd kwoty, 21 sprzeczność z
    Emmettem, 22 Zlecenie ze zdjęciami, 23 nieznane efenix z żółtym ``ADRES`` taty.

    Wiersz 19 ma starszą ``DATA``, więc proces kończy go przed wierszem 18.
    """

    red = Font(name="Arial", sz=8, color="FFFF0000", b=True, i=True)
    return [
        Order("#7701", day=date(2026, 9, 10), cells={"O": 1}, amount=_au({"012": 1}), quantities={"012": 1}, address_font=red),
        Order("#7702", day=date(2026, 9, 8), cells={"O": 1, "T": 2}, amount=_au({"012": 1, "019": 2}), address_font=None),
        Order("#7703", cells={"O": 1}, amount=_au({"012": 1}) + 5),
        Order("#7704", cells={"O": 1}, amount=_au({"012": 1}), emmett={"012": 2}),
        Order("#7705", cells={"AB": 1}),
        Order("#7706", address_fill=YELLOW),
    ]


def _marking_efenix(orders: Sequence[Order]) -> list[AtrapaOrder]:
    return [order.in_efenix(f"77{index}") for index, order in enumerate(orders) if order.number != "#7706"]


def test_znacznik_efenix_staje_tylko_w_adresie_gotowych_zlecen_a_reszta_pliku_zostaje_bajt_w_bajt(tmp_path: Path, page: Page) -> None:
    orders = _marking_orders()
    source = write_collective_file(tmp_path, orders)
    parts, looks, amounts = _parts(source), _address_looks(source), _amounts(source)
    with atrapa_efenix(_marking_efenix(orders)) as atrapa:
        outcome = _run(source, page, atrapa.url)

    assert (outcome.result, _figures(outcome)) == ("WYMAGA_SPRAWDZENIA", _figures_of(5, 2, warnings=1, photos=1))
    assert _address_fills(source) == {18: MARK_RGB, 19: MARK_RGB, 20: None, 21: None, 22: None, 23: YELLOW}
    # Poza wypełnieniem komórka ADRES ma swój wygląd, także czerwoną, pogrubioną kursywę taty w wierszu 18.
    assert _address_looks(source) == looks
    font = looks[18][0]
    assert (font.color.rgb if font.color else None, font.b, font.i) == ("FFFF0000", True, True)
    assert _amounts(source) == amounts
    written = _parts(source)
    assert [name for name, _ in written] == [name for name, _ in parts]
    before, after = dict(parts), dict(written)
    assert {name for name in before if before[name] != after[name]} == CHANGED_PARTS
    assert any(name.startswith("xl/externalLinks/") for name in before)
    sheet = "xl/worksheets/sheet1.xml"
    assert _without_address_styles(after[sheet], [18, 19]) == _without_address_styles(before[sheet], [18, 19])
    assert _without_appended_styles(after["xl/styles.xml"], before["xl/styles.xml"]) == before["xl/styles.xml"]
    # Jedno wypełnienie Znacznika i po jednym stylu na dwa dotychczasowe style gotowych ADRES.
    counts, original = _style_counts(after["xl/styles.xml"]), _style_counts(before["xl/styles.xml"])
    assert (counts[b"fills"] - original[b"fills"], counts[b"cellXfs"] - original[b"cellXfs"]) == (1, 2)


def test_nastepny_przebieg_na_pliku_ze_znacznikami_czyta_kwoty_au_ma_te_same_gotowe_i_niczego_nie_zmienia(
    tmp_path: Path, page: Page
) -> None:
    """Sprzeczność z Emmettem z pierwszego przebiegu zostawia w efenix Wycenę Emmetta, więc w drugim to „wycena już jest"."""

    orders = _marking_orders()
    source = write_collective_file(tmp_path, orders)
    with atrapa_efenix(_marking_efenix(orders)) as atrapa:
        _run(source, page, atrapa.url)
        marked, saved, state, requests = source.read_bytes(), source.stat().st_mtime_ns, atrapa.state(), len(atrapa.requested())
        second = _run(source, page, atrapa.url)
        after, again = atrapa.state(), atrapa.requested()[requests:]

    assert (_figures(second), _marked(source), _amounts(source)[18]) == (_figures_of(5, 2, photos=1), [18, 19], _au({"012": 1}))
    # Plik ze wszystkimi Znacznikami nie jest zapisywany od nowa, więc OneDrive nie dostaje nowej wersji.
    assert (source.read_bytes() == marked, source.stat().st_mtime_ns == saved) == (True, True)
    assert (after == state, set(again) <= READING) == (True, True)


def test_kolejne_uruchomienie_oznacza_nowe_gotowe_zlecenie_wypelnieniem_i_stylem_z_poprzedniego(
    tmp_path: Path, page: Page
) -> None:
    """Wiersz 19 jest gotowy dopiero w drugim przebiegu (w pierwszym efenix go nie zna), a jego ADRES ma ten sam styl co 18."""

    orders = [Order(f"#760{index}", cells={"O": 1}, amount=_au({"012": 1}), quantities={"012": 1}) for index in (1, 2)]
    source = write_collective_file(tmp_path, orders)
    original = _style_counts(dict(_parts(source))["xl/styles.xml"])
    with atrapa_efenix([orders[0].in_efenix("761")]) as atrapa:
        first = _run(source, page, atrapa.url)
    marked_first = _marked(source)
    with atrapa_efenix([order.in_efenix(f"76{index}") for index, order in enumerate(orders)]) as atrapa:
        second = _run(source, page, atrapa.url)

    assert (first.result, marked_first, second.result, _marked(source)) == ("WYMAGA_SPRAWDZENIA", [18], "OK", [18, 19])
    counts = _style_counts(dict(_parts(source))["xl/styles.xml"])
    assert (counts[b"fills"] - original[b"fills"], counts[b"cellXfs"] - original[b"cellXfs"]) == (1, 1)


def test_znacznik_efenix_staje_na_aktywnej_karcie_ktora_czyta_wycena_a_inny_arkusz_zostaje(tmp_path: Path, page: Page) -> None:
    order = Order("#7401", cells={"O": 1}, amount=_au({"012": 1}), quantities={"012": 1})
    source = write_collective_file(tmp_path, [order], notes_sheet_first=True)
    notes = dict(_parts(source))["xl/worksheets/sheet1.xml"]
    with atrapa_efenix([order.in_efenix("741")]) as atrapa:
        outcome = _run(source, page, atrapa.url)

    assert (outcome.result, _marked(source), dict(_parts(source))["xl/worksheets/sheet1.xml"] == notes) == ("OK", [18], True)


def test_gotowe_zlecenie_z_adresem_juz_w_kolorze_znacznika_zostaje_a_plik_sie_nie_zmienia(tmp_path: Path, page: Page) -> None:
    order = Order("#7501", cells={"O": 1}, amount=_au({"012": 1}), quantities={"012": 1}, address_fill=MARK_RGB)
    source = write_collective_file(tmp_path, [order])
    before = source.read_bytes()
    with atrapa_efenix([order.in_efenix("751")]) as atrapa:
        outcome = _run(source, page, atrapa.url)

    assert (outcome.result, _marked(source), source.read_bytes() == before) == ("OK", [18], True)


@contextmanager
def _excel_open(source: Path) -> Iterator[Callable[[], None]]:
    """Tata otwiera Plik zbiorczy w Excelu w trakcie przebiegu (plik blokady ``~$…``), a po przebiegu go zamyka."""

    lock = excel_lock_file(source)
    def open_in_excel() -> None:
        lock.write_bytes(b"")

    yield open_in_excel
    lock.unlink(missing_ok=True)


@contextmanager
def _saved_in_excel(source: Path) -> Iterator[Callable[[], None]]:
    """Tata zapisuje w trakcie przebiegu Plik zbiorczy ze zmienioną komórką."""

    def save() -> None:
        rewrite_parts(source, lambda name, data: data.replace(b">MIASTO<", b">MIASTA<") if name == "xl/worksheets/sheet1.xml" else data)

    yield save


@contextmanager
def _read_only_folder(source: Path) -> Iterator[Callable[[], None]]:
    """Folder okresu przestaje przyjmować nowe pliki w trakcie przebiegu (np. OneDrive go trzyma), a potem znów je przyjmuje."""

    folder = source.parent
    yield lambda: folder.chmod(0o555)
    folder.chmod(0o755)


@pytest.mark.parametrize(
    ("obstacle", "code"),
    [
        pytest.param(_excel_open, "EFENIX_ZNACZNIKI_PLIK_OTWARTY", id="otwarty-w-excelu"),
        pytest.param(_saved_in_excel, "EFENIX_ZNACZNIKI_PLIK_ZMIENIONY", id="zapisany-w-trakcie"),
        pytest.param(
            _read_only_folder,
            "EFENIX_ZNACZNIKI_NIEZAPISANE",
            id="zapis-nieudany",
            marks=pytest.mark.skipif(sys.platform == "win32", reason="chmod folderu nie odbiera zapisu na Windows"),
        ),
    ],
)
def test_plik_ktorego_nie_da_sie_oznaczyc_daje_komunikat_a_ponowne_uruchomienie_stawia_znaczniki_bez_zmian_w_efenix(
    tmp_path: Path, page: Page, obstacle: Callable[[Path], AbstractContextManager[Callable[[], None]]], code: str
) -> None:
    orders = [
        Order("#7801", cells={"O": 1}, amount=_au({"012": 1}), quantities={"012": 1}),
        Order("#7802", cells={"O": 1}, amount=_au({"012": 1})),
        Order("#7803", cells={"O": 1}, amount=_au({"012": 1}) + 5),
    ]
    source = write_collective_file(tmp_path, orders)
    with atrapa_efenix([order.in_efenix(f"78{index}") for index, order in enumerate(orders)]) as atrapa:
        with obstacle(source) as meanwhile:
            stopped = _run(source, page, atrapa.url, meanwhile=meanwhile)
            unmarked = source.read_bytes()
            assert (_stop(stopped), _marked(source)) == (code, [])
        priced, requests = atrapa.state(), len(atrapa.requested())
        assert source.read_bytes() == unmarked
        again = _run(source, page, atrapa.url)
        after, repeated = atrapa.state(), atrapa.requested()[requests:]

    assert {order.number: (order.quantities, order.execution_date, order.emmett_loads) for order in priced[0]} == {
        "7801": ({"012": 1}, None, 0),
        "7802": ({"012": 1}, DAY, 1),
        "7803": ({"012": 1}, DAY, 1),
    }
    assert (again.result, _figures(again), _marked(source)) == ("WYMAGA_SPRAWDZENIA", _figures_of(3, 2), [18, 19])
    assert (after == priced, set(repeated) <= READING) == (True, True)


@pytest.mark.parametrize(
    ("expire_after", "marked"),
    [
        pytest.param((SAVE_QUANTITY, 1), [18], id="po-pierwszym-gotowym"),
        pytest.param((SAVE_QUANTITY, 0), [], id="przed-pierwszym-gotowym"),
    ],
)
def test_zatrzymanie_w_polowie_przebiegu_oznacza_zlecenia_gotowe_przed_nim(
    tmp_path: Path, page: Page, expire_after: tuple[str, int], marked: list[int]
) -> None:
    orders = [Order(f"#790{index}", cells={"O": 1}, amount=_au({"012": 1})) for index in (1, 2, 3)]
    source = write_collective_file(tmp_path, orders)
    unmarked = source.read_bytes()
    with atrapa_efenix([order.in_efenix(f"79{index}") for index, order in enumerate(orders)], expire_after=expire_after) as atrapa:
        outcome = _run(source, page, atrapa.url)

    assert (_stop(outcome), _marked(source)) == ("EFENIX_WYLOGOWANY", marked)
    assert (source.read_bytes() == unmarked) == (not marked)


def test_plik_zbiorczy_otwarty_w_excelu_zatrzymuje_wycene_zanim_otworzy_efenix(tmp_path: Path, page: Page) -> None:
    orders = [Order("#7601", cells={"O": 1}, amount=_au({"012": 1}))]
    source = write_collective_file(tmp_path, orders)
    excel_lock_file(source).write_bytes(b"")
    with atrapa_efenix([orders[0].in_efenix("761")]) as atrapa:
        outcome = _run(source, page, atrapa.url)
        requested = atrapa.requested()

    assert (_stop(outcome), requested, _kept(outcome)) == ("EFENIX_PLIK_ZBIORCZY_OTWARTY", [], False)


def test_plik_zbiorczy_ze_znacznikami_daje_w_przebiegu_rozliczen_rozliczenie_pracownika_bez_nich(tmp_path: Path, page: Page) -> None:
    orders = _marking_orders()
    source = write_collective_file(tmp_path, orders)
    with atrapa_efenix(_marking_efenix(orders)) as atrapa:
        _run(source, page, atrapa.url)
    template = tmp_path / "Wzorzec rozliczenia pracownika.xlsx"
    write_template(template)
    assert _marked(source) == [18, 19]

    summary = run_settlements(source, write_mapping(tmp_path / "worker_mapping.yaml"), worker_template_path=template)

    assert summary.written_count == 1
    written = source.parent / f"Rozliczenie pracowników {PERIOD}" / f"Rozliczenie {PERIOD} - Adrian Testowy.xlsx"
    workbook = load_workbook(written)
    try:
        sheet = active_worksheet(workbook)
        fills = {row: sheet[f"B{row}"].fill for row in range(18, 24)}
        font = sheet["B18"].font
    finally:
        workbook.close()
    assert {row: (fill.fill_type, str(fill.fgColor.rgb)[-6:] if fill.fill_type else None) for row, fill in fills.items()} == {
        **{row: (None, None) for row in range(18, 23)},
        23: ("solid", YELLOW),
    }
    assert (font.color.rgb, font.b, font.i) == ("FFFF0000", True, True)
