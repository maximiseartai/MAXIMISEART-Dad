"""Wycena w efenix w testach: syntetyczny Plik zbiorczy z kwotami AU policzonymi „przez Excel” i zgodne z nim zlecenia
Atrapy efenix. Wszystkie numery, adresy i loginy są wymyślone.

Plik ma to, co zapis Znaczników efenix musi zostawić bajt w bajt, jak plik taty: wygląd komórek ``ADRES`` (czcionka,
obramowanie, ochrona, Wyróżnienia), szerokości kolumn, zewnętrzne łącze (``xl/externalLinks/*``) i wartości AU policzone przy
formule.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Protection, Side

from tests.fixtures.atrapa_efenix import AtrapaOrder
from tests.fixtures.settlement import COLLECTIVE_HEADERS, PERIOD, CellValue
from tests.fixtures.workbooks import active_worksheet, add_external_link_fixture

SOURCE_NAME = f"Rozliczenie {PERIOD} - zbiorcze.xlsx"
DAY = date(2026, 9, 9)
_THIN = Side(style="thin")


@dataclass(frozen=True)
class Order:
    """Zlecenie w obu miejscach naraz: Wiersz danych Pliku zbiorczego i to samo zlecenie w Atrapie efenix.

    ``cells``: dodatkowe komórki wiersza (litera → wartość), np. Ilości usług albo ilości w kolumnach Zlecenia ze
    zdjęciami. ``amount``: ``Kwota`` (AU), którą „policzył Excel”, domyślnie 0 jak w zleceniu bez Ilości usług; ``None``
    zostawia formułę bez policzonej wartości. ``quantities``: wycena, która już jest w efenix (kod → ilość), ``emmett``:
    propozycja Emmetta w tym samym kształcie. ``address_font`` i ``address_fill`` (RGB jednolitego wypełnienia) to
    Wyróżnienia komórki ``ADRES``, które tata nadał w Excelu; ``address_font=None`` zostawia ``ADRES`` bez żadnego stylu.
    """

    number: str
    worker: str = "adrian.testowy"
    address: str = "UL. TESTOWA 12/3"
    day: date | None = DAY
    cells: Mapping[str, CellValue] = field(default_factory=dict)
    amount: float | None = 0.0
    efenix_address: str = "Miasto-Dzielnica ul.Testowa 12 Lokal 3"
    quantities: Mapping[str, int] = field(default_factory=dict)
    emmett: Mapping[str, int] = field(default_factory=dict)
    address_font: Font | None = field(default_factory=lambda: Font(name="Arial", sz=8))
    address_fill: str | None = None

    def in_efenix(self, order_id: str) -> AtrapaOrder:
        """To zlecenie w Atrapie z identyfikatorem ``order_id``, zgodne z Plikiem zbiorczym, z jego wyceną i Emmettem."""

        return AtrapaOrder(
            order_id=order_id,
            number=self.number.removeprefix("#"),
            realizator=self.worker,
            address=self.efenix_address,
            quantities=dict(self.quantities),
            emmett=dict(self.emmett),
        )


def write_collective_file(directory: Path, orders: Sequence[Order], *, notes_sheet_first: bool = False) -> Path:
    """Plik zbiorczy okresu ``PERIOD`` w ``directory/PERIOD`` z nagłówkami Wzorca i Wierszem danych każdego zlecenia od 18.

    ``Kwota`` (AU) to formuła sumy z wartością policzoną przy „zapisie w Excelu”, jak w pliku taty. Komórka ``ADRES`` z
    czcionką ma też obramowanie i odblokowaną ochronę, jak w pliku taty, oraz Wyróżnienia zlecenia. ``notes_sheet_first``
    stawia przed arkuszem zleceń arkusz „Notatki” z tekstem w komórkach ``B`` tych samych wierszy, a arkusz zleceń zostaje
    aktywną kartą, którą czyta odczyt Pliku zbiorczego.
    """

    folder = directory / PERIOD
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / SOURCE_NAME
    workbook = Workbook()
    sheet = active_worksheet(workbook)
    sheet.title = "Sheet1"
    for column, header in enumerate(COLLECTIVE_HEADERS, start=1):
        sheet.cell(17, column).value = header
    sheet["AU17"] = "Kwota"
    sheet.column_dimensions["B"].width = 31.5
    sheet.column_dimensions["F"].width = 12.25
    computed: dict[int, float] = {}
    for number, order in enumerate(orders, start=18):
        moment = None if order.day is None else datetime(order.day.year, order.day.month, order.day.day, 10, 30)
        cells: dict[str, CellValue] = {
            "A": "MIASTO",
            "B": order.address,
            "F": order.number,
            "G": moment,
            "H": order.worker,
            "J": "tak",
            **order.cells,
        }
        for letter, value in cells.items():
            sheet[f"{letter}{number}"] = value
        address = sheet[f"B{number}"]
        if order.address_font is not None:
            address.font = order.address_font
            address.border = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
            address.protection = Protection(locked=False)
        if order.address_fill is not None:
            address.fill = PatternFill(fill_type="solid", fgColor=f"FF{order.address_fill}")
        sheet[f"AU{number}"] = f"=SUM(N{number}:AS{number})"
        if order.amount is not None:
            computed[number] = order.amount
    sheet_part = "xl/worksheets/sheet1.xml"
    if notes_sheet_first:
        notes = workbook.create_sheet("Notatki", 0)
        for number in range(17, 18 + len(orders)):
            notes[f"B{number}"] = "notatka"
        workbook.active = 1
        sheet_part = "xl/worksheets/sheet2.xml"
    workbook.save(path)
    add_external_link_fixture(path)
    _compute_amounts(path, computed, sheet_part)
    return path


def _compute_amounts(path: Path, amounts: Mapping[int, float], sheet_part: str) -> None:
    """Wpisuje do komórek AU wartość, którą Excel zapisuje przy formule (``<v>``), tak jak po zapisie w Excelu."""

    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts[sheet_part].decode("utf-8")
    for row, amount in amounts.items():
        sheet, count = re.subn(
            rf'(<c r="AU{row}"[^>]*><f>[^<]*</f>)<v ?/>', rf"\g<1><v>{amount}</v>", sheet
        )
        assert count == 1, f"AU{row}"
    parts[sheet_part] = sheet.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
