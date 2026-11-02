"""Zakres Wyceny w efenix: zlecenia Pliku zbiorczego, które proces sprawdza w efenix, w kolejności taty.

Jedyną operacją jest ``read_scope()``: jeden odczyt Pliku zbiorczego z wartościami, które policzył Excel
(``workbook_files.read_computed_rows``), i czysta reguła zakresu (``scope``). W zakresie jest każdy Wiersz danych z
Wykonawcą, który nie jest Zleceniem ze zdjęciami (AB, AC, AD, AE, AF, AJ–AP, AS) ani usterką (AI); wartość 0 i pusta
komórka znaczą to samo, a każda inna treść, także tekst, czyni wiersz Zleceniem ze zdjęciami albo usterką. Wiersz, który
jest jednym i drugim, liczy się jako Zlecenie ze zdjęciami. Kolejność to Wykonawcy alfabetycznie, a w każdym ``DATA`` od
najstarszej (wiersz bez daty na końcu Wykonawcy, remis po numerze wiersza). Plik zbiorczy jest tylko czytany, a zakres
niesie jego bajty z tego odczytu i kolumnę ``ADRES``, w których Wycena w efenix stawia Znaczniki efenix (``efenix_marking``).
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Final

from openpyxl.utils import column_index_from_string, get_column_letter

from . import workbook_files
from .order_facts import TataValue, tata_value
from .preparation_register import ADDRESS, DATE, WORKER, Column
from .workbook_files import ComputedRow, cell_datetime, is_blank
from .worker_config import normalize_text

__all__ = ["AmountsNotComputed", "CollectiveOrder", "MissingColumns", "PricingScope", "read_scope", "scope"]

_NUMBER: Final = Column("NUMER ZL.", "NUMER ZL.")
_COLUMNS: Final = (_NUMBER, ADDRESS, DATE, WORKER)


def _letters(first: str, last: str) -> tuple[str, ...]:
    return tuple(get_column_letter(column) for column in range(column_index_from_string(first), column_index_from_string(last) + 1))


PHOTO_COLUMNS: Final = ("AB", "AC", "AD", "AE", "AF", *_letters("AJ", "AP"), "AS")
"""Kolumny Zlecenia ze zdjęciami (GLOSSARY)."""
DEFECT_COLUMN: Final = "AI"
QUANTITY_COLUMNS: Final = (*_letters("N", "AP"), "AR", "AS")
"""Kolumny Ilości usług, które spec #741 przypisuje pozycjom efenix albo wymienia jako kolumny bez pary (N:AP, AR, AS)."""
AMOUNT_COLUMN: Final = "AU"


@dataclass(frozen=True)
class CollectiveOrder:
    """Zlecenie w zakresie Wyceny w efenix: Wiersz danych Pliku zbiorczego w tym, co proces porównuje z efenix.

    ``number``: ``NUMER ZL.`` bez ``#`` i spacji, ``None`` przy pustej komórce. ``day``: dzień z ``DATA``, ``None``, gdy
    komórka nie jest datą. ``quantities``: litera kolumny Ilości usług → ilość taty różna od 0. ``amount``: ``Kwota`` (AU)
    policzona przez Excel.
    """

    row: int
    number: str | None
    worker: str
    address: str
    day: date | None
    quantities: Mapping[str, TataValue]
    amount: Decimal


@dataclass(frozen=True)
class PricingScope:
    """Zlecenia w zakresie w kolejności, w której proces je sprawdza, i liczby pominiętych Wierszy danych z Wykonawcą.

    ``address_column``: litera kolumny ``ADRES``. ``content``: bajty Pliku zbiorczego, z których pochodzi zakres.
    """

    orders: tuple[CollectiveOrder, ...]
    with_photos: int
    defects: int
    address_column: str
    content: bytes


@dataclass(frozen=True)
class AmountsNotComputed:
    """Wiersz w zakresie nie ma ``Kwota`` (AU) policzonej przez Excel; ``row`` to pierwszy taki wiersz."""

    row: int


@dataclass(frozen=True)
class MissingColumns:
    """W wierszu 17 Pliku zbiorczego nie ma nagłówka kolumny, którą proces porównuje z efenix."""

    headers: tuple[str, ...]


def read_scope(path: Path) -> PricingScope | AmountsNotComputed | MissingColumns:
    """Zakres Wyceny w efenix z Pliku zbiorczego ``path``; ``WorkbookError``, gdy pliku nie da się odczytać albo jest
    otwarty w Excelu."""

    read = workbook_files.read_computed_rows(path)
    return scope(read.header, read.rows, read.content)


def scope(
    header: Mapping[str, int], rows: Iterable[ComputedRow], content: bytes
) -> PricingScope | AmountsNotComputed | MissingColumns:
    """Zakres z Wierszy danych w kolejności z pliku o bajtach ``content``; ``header``: znormalizowany nagłówek wiersza 17
    → kolumna od 1."""

    names = [normalize_text(column.target_header) for column in _COLUMNS]
    missing = tuple(column.target_header for column, name in zip(_COLUMNS, names, strict=True) if name not in header)
    if missing:
        return MissingColumns(missing)
    number, address, started, worker = (get_column_letter(header[name]) for name in names)
    orders: list[CollectiveOrder] = []
    with_photos = defects = 0
    for row in rows:
        login = _text(row.value(worker))
        if not login:
            continue
        if any(_filled(row.value(column)) for column in PHOTO_COLUMNS):
            with_photos += 1
            continue
        if _filled(row.value(DEFECT_COLUMN)):
            defects += 1
            continue
        amount = _amount(row.value(AMOUNT_COLUMN))
        if amount is None:
            return AmountsNotComputed(row.number)
        moment = cell_datetime(row.value(started))
        orders.append(
            CollectiveOrder(
                row=row.number,
                number=_order_number(row.value(number)),
                worker=login,
                address=_text(row.value(address)),
                day=None if moment is None else moment.date(),
                quantities=_quantities(row),
                amount=amount,
            )
        )
    return PricingScope(tuple(sorted(orders, key=_tata_order)), with_photos, defects, address, content)


def _tata_order(order: CollectiveOrder) -> tuple[str, bool, date, int]:
    return (order.worker.casefold(), order.day is None, order.day or date.min, order.row)


def _filled(value: object) -> bool:
    return tata_value(value) != 0


def _text(value: object) -> str:
    return "" if is_blank(value) else " ".join(str(value).split())


def _order_number(value: object) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = "".join(_text(value).removeprefix("#").split())
    return text or None


def _amount(value: object) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    amount = Decimal(str(value))
    return amount if amount.is_finite() else None


def _quantities(row: ComputedRow) -> dict[str, TataValue]:
    values = {letter: tata_value(row.value(letter)) for letter in QUANTITY_COLUMNS}
    return {letter: value for letter, value in values.items() if value != 0}
